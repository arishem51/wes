#!/usr/bin/env python3
"""Sync the official SEP490_G4 report docs from Google Drive into a greppable
text mirror under wes/report/specs/.

    python wes/scripts/sync-specs.py            # download + extract all
    python wes/scripts/sync-specs.py 3 4        # only Report 3 and 4
    python wes/scripts/sync-specs.py --local    # re-extract cached downloads

Drive folder: https://drive.google.com/drive/folders/1UjrNCm58OVG_p-GDwhyF7tYx8nt_qrJH
"""

import json
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
import docx
import openpyxl

FOLDER_ID = "1UjrNCm58OVG_p-GDwhyF7tYx8nt_qrJH"

DOCS = [
    ("1", "15e5H0MC_LfcBD4qQTn6f8MCwiF4-Brvh", "docx", "report1-vision-scope"),
    ("2", "1d2Ev9cqIF0zMaFyKOEoKuGJPjx6mPFXV", "docx", "report2-project-plan"),
    ("2.1", "11903Kd_Ro40B4R4aytKN1Fhvokd8LE2z", "xlsx", "report2.1-project-tracking"),
    ("3", "1sLImCWyhR1zsPeqJRALl5cyBdc9ioGgo", "docx", "report3-srs"),
    ("4", "1CdQtuqG_uXj1p39vkQ-2SEOGb1zfcWet", "docx", "report4-sds"),
    ("5.0", "1KUw-qlsLUqn5Uf6aXB7hSF4Ojcr4pbJt", "docx", "report5.0-test-documentation"),
    ("5.2", "1ZZEzImpKNF-K2f5fzOuNCAeIz3AqidT-", "xlsx", "report5.2-integration-test"),
    ("5.3", "1LYABMblvFA9hTXYjTAqLYleiFNsaZmwZ", "xlsx", "report5.3-system-test"),
    ("weekly", "1Kd4m0Q9J_QCReyGdkE92lGCMCxCqgtzt", "xlsx", "weekly-report"),
]

REPO = Path(__file__).resolve().parents[1]
OUT_DIR = REPO / "report" / "specs"
CACHE_DIR = OUT_DIR / ".cache"


def download(file_id: str, dest: Path) -> None:
    url = f"https://drive.google.com/uc?export=download&id={file_id}"
    subprocess.run(
        ["curl", "-sSL", "-o", str(dest), url],
        check=True,
    )
    if dest.stat().st_size < 4096:
        raise RuntimeError(f"{dest.name}: download too small, likely an HTML error page")


def iter_blocks(elm, parent):
    for child in elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)
        elif child.tag == qn("w:sdt"):
            for content in child.iterchildren(qn("w:sdtContent")):
                yield from iter_blocks(content, parent)


def heading_level(paragraph: Paragraph) -> int:
    name = (paragraph.style.name or "") if paragraph.style else ""
    if name.startswith("Heading"):
        tail = name.replace("Heading", "").strip()
        return int(tail) if tail.isdigit() else 1
    if name in ("Title",):
        return 1
    return 0


def cell_text(cell: _Cell) -> str:
    parts = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    return " / ".join(parts).replace("|", "\\|")


def table_to_md(table: Table) -> list[str]:
    rows = []
    for row in table.rows:
        rows.append([cell_text(c) for c in row.cells])
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    for r in rows[1:]:
        out.append("| " + " | ".join(r) + " |")
    return out


def extract_docx(path: Path, title: str, meta: dict) -> str:
    doc = docx.Document(str(path))
    lines = [f"# {title}", ""]
    lines += front_matter(meta)
    for block in iter_blocks(doc.element.body, doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            level = heading_level(block)
            lines.append(f"{'#' * min(level + 1, 6)} {text}" if level else text)
            lines.append("")
        else:
            md = table_to_md(block)
            if md:
                lines += md + [""]
    return "\n".join(lines).rstrip() + "\n"


def extract_xlsx(path: Path, title: str, meta: dict) -> str:
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    lines = [f"# {title}", ""]
    lines += front_matter(meta)
    for ws in wb.worksheets:
        lines += [f"## Sheet: {ws.title}", ""]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip().replace("|", "\\|") for v in row]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                rows.append(cells)
        if not rows:
            lines += ["_(empty)_", ""]
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        lines.append("| " + " | ".join(rows[0]) + " |")
        lines.append("|" + "---|" * width)
        for r in rows[1:]:
            lines.append("| " + " | ".join(r) + " |")
        lines.append("")
    wb.close()
    return "\n".join(lines).rstrip() + "\n"


def front_matter(meta: dict) -> list[str]:
    return [
        "> GENERATED FILE — do not edit by hand.",
        f"> Source: Google Drive `{meta['file_id']}` (folder {FOLDER_ID})",
        f"> Synced: {meta['synced_at']}",
        "> Regenerate: `python wes/scripts/sync-specs.py`",
        "",
        "---",
        "",
    ]


def main() -> int:
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    local_only = "--local" in sys.argv
    selected = [d for d in DOCS if not args or d[0] in args]
    if not selected:
        print(f"no doc matched {args}; known: {[d[0] for d in DOCS]}")
        return 1

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    CACHE_DIR.mkdir(parents=True, exist_ok=True)

    manifest_path = OUT_DIR / "_manifest.json"
    manifest = json.loads(manifest_path.read_text()) if manifest_path.exists() else {}

    for key, file_id, kind, slug in selected:
        raw = CACHE_DIR / f"{slug}.{kind}"
        if not local_only:
            print(f"[{key}] downloading {slug}.{kind} ...")
            download(file_id, raw)
        elif not raw.exists():
            print(f"[{key}] SKIP — no cached download at {raw}")
            continue

        meta = {
            "file_id": file_id,
            "synced_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "bytes": raw.stat().st_size,
        }
        title = f"Report {key} — {slug.split('-', 1)[1].replace('-', ' ')}"
        body = extract_docx(raw, title, meta) if kind == "docx" else extract_xlsx(raw, title, meta)
        out = OUT_DIR / f"{slug}.md"
        out.write_text(body, encoding="utf-8")
        manifest[slug] = meta
        print(f"[{key}] -> {out.relative_to(REPO)}  ({len(body.splitlines())} lines)")

    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(f"\nmanifest: {manifest_path.relative_to(REPO)}")
    print("NOTE: Report 5.1 Unit Test is legacy .xls — not mirrored (openpyxl cannot read it).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
